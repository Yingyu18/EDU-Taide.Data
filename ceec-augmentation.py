import os
from anthropic import Anthropic 
from docx import Document
from dotenv import load_dotenv

# 加載 .env 文件中的變數
load_dotenv()

# 使用環境變數
api_key = os.getenv("Anthropic_key")

def read_docx(file_path):
    # 讀取 docx 檔案內容 
    try:
        doc = Document(file_path)      
        full_text = []
        
        題目=""
        答題=""
        評分=""
        question=True
        response=False
        # 段落
        for para in doc.paragraphs:
            if "題目" in para.text:
                continue
            if "答題" in para.text:
                question=False
                continue
            if "評分" in para.text:
                評分=para.text.strip("評分： ")
                if 評分=="16" or 評分=="17":
                    評分="B+"
                elif 評分=="22" or 評分=="23":
                    評分="A+"
                else:
                    評分="A"
                break
            if question==True:
                題目+=para.text
            else:
                答題+=para.text

            if para.text.strip():  # 排除空白段落
                full_text.append(para.text)
        
        res = { 
            "題目":題目,
            "答題":答題,
            "評分":評分
        }
        print(res)
        return res
        
    except Exception as e:
        print(f"讀取 docx 檔案時發生錯誤: {e}")
        return None

def process_multiple_files():
    # years = ["108", "109", "110"]
    # nos = ["1", "2", "3", "4", "5"]
    # years = ["107"]
    # years = ["113","112", "111", "108", "109", "110"]
    years = ["111"]
    nos = ["1"]
    
    for year in years:
        for no in nos:
            # file_path = f"./file/{year}/ceec-{year}-{no}.docx"
            file_path = f"./file/{year}/ceec-{year}-{no}/ceec-{year}-{no}.docx"
            
            if os.path.exists(file_path):
                # 讀取文件內容
                content = read_docx(file_path)
                if content:
                    # 處理文件內容
                    # TODO: iterate different level article inside the prompt 
                    levels = ["B", "B+", "C", "C+"]
                    # levels = ["C+"]
                    for level in levels:
                        process_content(content, level, year, no)
            else:
                print(f"檔案不存在: {file_path}")

def process_content(content,level, year, no): # content is a json object 
    """處理文件內容"""
    try:
        client = Anthropic(api_key=api_key)
        # Replace placeholders like {{題目}} with real values,
        # because the SDK does not support variables.
        題目=content["題目"]
        答題=content["答題"]
        原等級=content["評分"]
        改寫等級=level
        print(題目,"\n", 答題,"\n", 原等級, "\n")
        
        response = client.messages.create(
            model="claude-3-5-sonnet-20241022",
            max_tokens=1000,
            temperature=0,
            system="你現在是一位文字工作者，任務是要改寫文章。請依據給定的評分標準和答題，調整答題文章在評分原則中五個向度的水準，將原本級的文章改寫為改寫等級的文章。輸出改寫文章，並為改寫文章輸出「整體分數」和「寫作維度」分數。請勿輸出其他內容。\n\n# 評分標準\n\n- 寫作維度（每個維度0-5分）：\n  - 立意：評估主題或論點是否清晰、深刻。\n  - 結構：檢查文章是否組織有序、結構嚴謹。\n  - 修辭：語言的運用是否恰當、有效。\n  - 敘述：評估故事或論述的發展是否充分、合理。\n  - 啟發：核查文章是否具有啟發性，能否引發讀者思考或情感共鳴。\n\n- 整體分數（滿分25分）：\n  - A+（22-25分）：立意精切，結構嚴謹，條理清晰，文辭精練，整體具有啟發性或感動人心的力量。\n  - A（18-21分）：立意深刻，結構合理，敘述有序，文辭得當，具反省或感發效果。\n  - B+（14-17分）：立意適當，結構尚可，表達清楚，文辭順暢，觀點平穩，情感表達穩妥。\n  - B（10-13分）：立意尚可，結構稍具，表述大致合宜，觀點屬習見或情感表達較為平淡。\n  - C+（6-9分）：立意較為平淺，表達有限，文章分段隨意，敘述空泛，文辭平凡，觀點粗糙或情意淡薄。\n  - C（1-5分）：立論單薄或空泛，文章結構混亂，敘述邏輯不清，文辭簡陋，觀點錯誤或缺乏情感。\n  - 0分：空白卷，文不對題，或僅抄錄題幹。\n\n- 其他考量：\n  - 若作文為只有一個段落，最高只能給B+（14-17分）。\n  - 「*」符號表示錯字。若錯字過多，酌情扣分。\n  - 若未書寫要求的內容，酌情扣分。\n\n# Output Format\n寫作維度: 立意: X/5, 結構: X/5, 修辭: X/5, 敘述: X/5, 啟發: X/5\n總分: XX/25\n\n# Steps\n1. 根據上述每個「寫作維度」對答題作文不同維度進行等級調整，改寫為另一篇文章。\n2. 對改寫文章進行詳細評估，給出「整體分數」，以及各維度分數。\n3. 嚴格依照 Output Format 輸出分數。請勿輸出其他內容。",
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": "<examples>\n<example>\n<答題>\n在網路迅速發展的時代，人類的記憶力到底是會衰退，還是隨著科技愈發博學多聞？我想這必須考慮到：是「人使用網路」抑或是「人依賴網路」？一旦網路消失了，人們會再創顛峰，又或是自甘墮落？\n\n我對網路的發展是抱持著正向的想法的，也許有人認為這不利於認知學習，但對於一個極需在其專業領域上發揮創造力的人而言，若是自小被教育了許多過去人認為是對的「知識」，那他要如何跳脫框架思考？正因為「網路」的言論充滿了不確定性，才促使一個人產生懷疑而充滿好奇，進而做更深入的研究。正因為懷疑，哥白尼才提出日心說，才顛覆了千年來人們的信念。\n\n對一個科學家或是發明家來說，毋須記憶大量知識才更能顛覆固有思考，進而專精在自己所學，並且在使用網路時記憶下對研究有助的學說，也許哪一天也能再創科學界的顛峰！這樣的思維模式使得人們能夠從現有的資訊中抽取關鍵，並且將其與自己的知識相結合，從而形成更具創意的觀點和解決方案。以此而言，網路不僅是知識的載體，更是創新思維的重要工具。\n\n然而，若完全依賴網路，則可能導致思考的淺薄和記憶的衰退。對於學習而言，記憶的過程仍是關鍵的一環。人類的記憶力在某種程度上與對知識的理解和應用能力息息相關，因此若能在網路的輔助下，兼顧自主學習和知識的儲存，或許能讓我們在科技發展的浪潮中更加茁壯。\n\n此外，網路的發展也要求教育者重新思考教學的方式。傳統的知識傳授模式應該轉變為培養學生的批判性思維和解決問題的能力。鼓勵學生在探索網路資訊的同時，發展出獨立思考的能力，這樣才能在未來的社會中適應不斷變化的環境。\n\n總結來說，網路的發展無疑改變了我們獲取和處理資訊的方式，關鍵在於如何使用這一工具。透過合理的方式運用網路，將有助於提升人類的創造力與學習能力；反之，過度依賴則可能使我們逐漸失去思考的能力。因此，在面對科技進步的同時，我們更應該思考如何平衡網路使用與自主學習的關係，才能讓記憶力與思考力持續增長。\n\n</答題>\n<題目>\n自從有了電腦、智慧型手機及網路搜尋引擎之後，資訊科技的發展改變了人類大腦處理資訊的方式。我們可能儲存了大量的資訊，卻來不及閱讀，也不再費力記憶周遭事物和相關知識，因為只要輕鬆點一下滑鼠、滑一下手機，資訊就傳到我們面前。\n二十一世紀資訊量以驚人的速度暴增，有人認為網路資訊易於取得，會使記憶力與思考力衰退，不利於認知學習；也有人視網際網路為人類的外接大腦記憶體，意味著我們無須記憶大量知識，而可以專注在更重要、更有創造力的事物上。對於以上兩種不同的觀點，請提出你個人的看法。文長限700字內。\n\n</題目>\n<原等級>\nA\n</原等級>\n<改寫等級>\nC\n</改寫等級>\n<ideal_output>\n現在的人都在用網路，每天都在滑手機，很少人會想要去記東西。我覺得這樣不太好，因為人應該要多動腦想事情，不能都靠網路。\n網路雖然可以讓我們找到很多資料，但是我覺得還是要自己記東西比較好。像我以前都很認真讀書，自己記筆記，成績就很好。現在的同學都只會用手機，什麼都不記，上課也不認真聽，都在玩手機，這樣很浪費時間。\n網路也不是說有多好，因為網路上有很多錯誤的資料。而且網路會讓人變笨，因為大家都不想要思考，看到什麼都直接相信。我認為人還是要自己動腦想，不能完全靠網路。\n很多人說網路很好用，可以讓我們不用記那麼多東西。但是我覺得這樣不行，因為如果都沒有記憶，人的大腦就會變成空的，什麼都不會。而且網路有時候會壞掉，如果那時候找不到資料怎麼辦？\n所以我希望大家可以少用一點網路，多用腦袋想事情。網路應該只是一個輔助工具，不能變成我們的大腦。我們要記住，人比網路還要厲害！\n\n</ideal_output>\n</example>\n</examples>\n\n"
                        },
                        {
                            "type": "text",
                            "text": f"### 題目：{題目}\n\n\n\n### 答題：{答題}\n\n上述文章為評分標準中 {原等級} 的文章，請將此篇文章改寫為{改寫等級}級文章，並從五個寫作維度替此篇 {改寫等級} 級文章評分。"
                        }
                    ]
                }
            ]
        )
        
        resContent = response.content[0].text
        print(response.content[0].text)

        # # 建立檔案名稱
        file_name = f"ceec-{year}-{no}-{level}.txt"
        
        # # 確保目標資料夾存在
        output_dir = f"./file/{year}/ceec-{year}-{no}"
        os.makedirs(output_dir, exist_ok=True)
        
        # # 完整檔案路徑
        file_path = os.path.join(output_dir, file_name)
        
        
        # 將結果寫入檔案
        with open(file_path, 'w', encoding='utf-8') as f:
            # 寫入基本資訊
            f.write(resContent)
        
        print(f"已儲存結果至：{file_path}")
        
    except Exception as e:
        print(f"處理內容時發生錯誤: {e}")

def main():
    # 首先安裝必要套件
    # pip install python-docx anthropic
    
    process_multiple_files()

if __name__ == "__main__":
    main()